"""Write extracted invoices to a formatted, analyst-ready Excel workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from invoice_extractor.extractor import Invoice

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
MONEY = "#,##0.00"
DATE = "yyyy-mm-dd"


def _write_table(ws: Worksheet, headers: list[str], rows: list[list], money_cols=(), date_cols=()) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = MONEY
    for col in date_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = DATE
    for idx, header in enumerate(headers, start=1):
        values = [len(str(c.value)) for c in ws[get_column_letter(idx)] if c.value is not None]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(values + [len(header)]) + 3, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(invoices: list[Invoice], output: str | Path) -> Path:
    """Create the workbook with Invoices, Line Items, Review and Summary sheets."""
    output = Path(output)
    wb = Workbook()

    ws = wb.active
    ws.title = "Invoices"
    rows = [[inv.source_file, inv.vendor, inv.tax_id, inv.number, inv.date, inv.due_date, inv.currency,
             inv.subtotal, inv.tax, inv.total, "OK" if inv.is_valid else "REVIEW"] for inv in invoices]
    _write_table(ws, ["File", "Vendor", "Tax ID", "Invoice #", "Date", "Due date", "Currency",
                      "Subtotal", "Tax", "Total", "Status"], rows, money_cols=(8, 9, 10), date_cols=(5, 6))
    for row in ws.iter_rows(min_row=2):
        if row[-1].value == "REVIEW":
            for cell in row:
                cell.fill = ERROR_FILL

    items = wb.create_sheet("Line Items")
    item_rows = [[inv.number, inv.vendor, it.description, it.quantity, it.unit_price, it.amount]
                 for inv in invoices for it in inv.items]
    _write_table(items, ["Invoice #", "Vendor", "Description", "Qty", "Unit price", "Amount"], item_rows,
                 money_cols=(5, 6))

    review = wb.create_sheet("Review")
    issue_rows = [[inv.source_file, inv.number, issue] for inv in invoices for issue in inv.issues]
    _write_table(review, ["File", "Invoice #", "Issue"], issue_rows or [["-", "-", "No issues found"]])

    summary = wb.create_sheet("Summary", 0)
    summary.append(["Invoice extraction summary"])
    summary["A1"].font = Font(bold=True, size=14)
    last = len(invoices) + 1
    summary.append(["Invoices processed", len(invoices)])
    summary.append(["Invoices needing review", f'=COUNTIF(Invoices!K2:K{last},"REVIEW")'])
    summary.append(["Line items extracted", len(item_rows)])
    summary.append([])
    summary.append(["Currency", "Total invoiced"])
    for cell in summary[6]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for currency in sorted({inv.currency for inv in invoices if inv.currency}):
        summary.append([currency, f'=SUMIF(Invoices!G2:G{last},"{currency}",Invoices!J2:J{last})'])
        summary.cell(summary.max_row, 2).number_format = MONEY
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
