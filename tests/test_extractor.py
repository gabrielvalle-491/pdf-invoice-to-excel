from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from invoice_extractor.excel_writer import write_workbook
from invoice_extractor.extractor import Invoice, LineItem, extract_folder, validate
from invoice_extractor.generate_samples import generate
from invoice_extractor.parsing import parse_amount, parse_date


@pytest.mark.parametrize("raw, expected", [
    ("$ 1,234.56", Decimal("1234.56")),
    ("ARS 1.234,56", Decimal("1234.56")),
    ("1234,5", Decimal("1234.50")),
    ("1,234", Decimal("1234.00")),
    ("1.234.567", Decimal("1234567.00")),
    ("USD 0.99", Decimal("0.99")),
    ("", None),
    ("n/a", None),
])
def test_parse_amount(raw, expected):
    assert parse_amount(raw) == expected


@pytest.mark.parametrize("raw", ["15/03/2026", "2026-03-15", "15-03-2026", "Mar 15, 2026"])
def test_parse_date_formats(raw):
    assert parse_date(raw) == date(2026, 3, 15)


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    folder = tmp_path_factory.mktemp("invoices")
    truth = generate(folder, count=8)
    return folder, {t["file"]: t for t in truth}


def test_extracts_every_field_from_generated_invoices(samples):
    folder, truth = samples
    invoices = extract_folder(folder)
    assert len(invoices) == len(truth)
    for inv in invoices:
        expected = truth[inv.source_file]
        assert inv.number == expected["number"]
        assert inv.vendor == expected["vendor"]
        assert inv.date == expected["date"]
        assert inv.subtotal == expected["subtotal"]
        assert inv.tax == expected["tax"]
        assert inv.total == expected["total"]
        assert len(inv.items) == expected["items"]
        assert inv.is_valid, inv.issues


def test_validation_flags_inconsistent_totals():
    inv = Invoice(source_file="x.pdf", number="1", date=date(2026, 1, 1), subtotal=Decimal("100"),
                  tax=Decimal("21"), total=Decimal("150"),
                  items=[LineItem("A", Decimal("1"), Decimal("90"), Decimal("90"))])
    issues = validate(inv)
    assert any("subtotal" in i for i in issues)
    assert any("does not match total" in i for i in issues)


def test_corrupted_pdf_goes_to_review(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    [inv] = extract_folder(tmp_path)
    assert not inv.is_valid


def test_workbook_has_all_sheets(samples, tmp_path):
    folder, _ = samples
    path = write_workbook(extract_folder(folder), tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Invoices", "Line Items", "Review"]
    assert wb["Invoices"].max_row == 9
